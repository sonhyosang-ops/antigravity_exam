import io
import openpyxl
from report_generator import ExcelReportGenerator

def test_report_generation():
    # 1. 가상의 시험 기준표 데이터 (소문항 포함)
    mock_criteria = {
        "exam_title": "2026학년도 1학기 컴퓨터응용 시험",
        "total_score": 100.0,
        "items": [
            {"item_no": "1-1", "score": 3.0, "type": "객관식", "correct_answer": "4"},
            {"item_no": "1-2", "score": 3.0, "type": "단답형", "correct_answer": "CPU"},
            {"item_no": "2", "score": 4.0, "type": "객관식", "correct_answer": "2"},
            {"item_no": "3", "score": 5.0, "type": "서술형", "correct_answer": "메모리와 연산장치 제어"},
            {"item_no": "6-1", "score": 5.0, "type": "단답형", "correct_answer": "ALU"},
            {"item_no": "6-2", "score": 5.0, "type": "서술형", "correct_answer": "명령어 레지스터를 통한 디코딩"}
        ]
    }
    
    # 2. 서로 다른 2개 학과의 학생 데이터
    mock_students = [
        {
            "department": "기계공학과",
            "student_id": "20240101",
            "name": "홍길동",
            "results": {
                "1-1": {"score": 3.0, "status": "O", "answer": "4"},
                "1-2": {"score": 3.0, "status": "O", "answer": "CPU"},
                "2": {"score": 4.0, "status": "O", "answer": "2"},
                "3": {"score": 4.0, "status": "△", "answer": "연산장치 제어"},
                "6-1": {"score": 5.0, "status": "O", "answer": "ALU"},
                "6-2": {"score": 3.0, "status": "△", "answer": "명령어 디코딩"}
            },
            "total_score": 22.0
        },
        {
            "department": "기계공학과",
            "student_id": "20240102",
            "name": "김철수",
            "results": {
                "1-1": {"score": 0.0, "status": "X", "answer": "1"},
                "1-2": {"score": 3.0, "status": "O", "answer": "CPU"},
                "2": {"score": 4.0, "status": "O", "answer": "2"},
                "3": {"score": 0.0, "status": "X", "answer": "오답"},
                "6-1": {"score": 0.0, "status": "X", "answer": "오답"},
                "6-2": {"score": 0.0, "status": "X", "answer": "오답"}
            },
            "total_score": 7.0
        },
        {
            "department": "스마트소프트웨어과",
            "student_id": "20240201",
            "name": "이영희",
            "results": {
                "1-1": {"score": 3.0, "status": "O", "answer": "4"},
                "1-2": {"score": 3.0, "status": "O", "answer": "CPU"},
                "2": {"score": 4.0, "status": "O", "answer": "2"},
                "3": {"score": 5.0, "status": "O", "answer": "메모리와 연산장치 제어"},
                "6-1": {"score": 5.0, "status": "O", "answer": "ALU"},
                "6-2": {"score": 5.0, "status": "O", "answer": "명령어 레지스터를 통한 디코딩"}
            },
            "total_score": 25.0
        }
    ]
    
    # 3. 엑셀 생성
    excel_bytes = ExcelReportGenerator.generate_excel_report(mock_criteria, mock_students)
    assert len(excel_bytes) > 0, "엑셀 바이너리가 비어있습니다."
    
    # 4. 생성된 엑셀 검증
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    print(f"생성된 시트 목록: {sheet_names}")
    
    # 학과별 시트 자동 추가 확인
    assert "기계공학과" in sheet_names, "'기계공학과' 시트가 존재해야 합니다."
    assert "스마트소프트웨어과" in sheet_names, "'스마트소프트웨어과' 시트가 존재해야 합니다."
    
    # 기계공학과 시트 검증
    ws = wb["기계공학과"]
    
    # 1행: [공란, 문항, 1-1, 1-2, 2, 3, 6-1, 6-2, 득점, 순위]
    row1 = [ws.cell(1, col).value for col in range(1, 11)]
    print(f"기계공학과 1행: {row1}")
    assert row1[1] == "문항"
    assert row1[2] == "1-1"
    assert row1[8] == "득점"
    assert row1[9] == "순위"
    
    # 2행: [공란, 정답, 4, CPU, 2, ..., 총배점, 공란]
    row2 = [ws.cell(2, col).value for col in range(1, 11)]
    print(f"기계공학과 2행: {row2}")
    assert row2[1] == "정답"
    assert row2[2] == "4"
    assert row2[3] == "CPU"
    
    # 3행: [학번, 성명, ...]
    row3 = [ws.cell(3, col).value for col in range(1, 11)]
    print(f"기계공학과 3행: {row3}")
    assert row3[0] == "학번"
    assert row3[1] == "성명"
    
    # 4행: 학생 1 (홍길동)
    row4 = [ws.cell(4, col).value for col in range(1, 11)]
    print(f"기계공학과 4행 (학생1): {row4}")
    assert row4[0] == "20240101"
    assert row4[1] == "홍길동"
    assert row4[8] == "=SUM(C4:H4)"
    assert row4[9] == "=RANK(I4, $I$4:$I$5)"
    
    # 5행: 학생 2 (김철수)
    row5 = [ws.cell(5, col).value for col in range(1, 11)]
    print(f"기계공학과 5행 (학생2): {row5}")
    assert row5[0] == "20240102"
    assert row5[1] == "김철수"
    assert row5[8] == "=SUM(C5:H5)"
    assert row5[9] == "=RANK(I5, $I$4:$I$5)"
    
    # 스마트소프트웨어과 시트 검증
    ws_sw = wb["스마트소프트웨어과"]
    row4_sw = [ws_sw.cell(4, col).value for col in range(1, 11)]
    print(f"소프트웨어과 4행 (학생1): {row4_sw}")
    assert row4_sw[0] == "20240201"
    assert row4_sw[1] == "이영희"
    
    print("\n✅ 모든 구글 스프레드시트 규격 및 학과별 자동 시트 분할 테스트 통과!")

if __name__ == "__main__":
    test_report_generation()
