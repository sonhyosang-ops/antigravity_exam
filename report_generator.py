import io
import re
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelReportGenerator:
    """구글 스프레드시트 양식과 100% 동일한 학과별 자동 시트 분할 엑셀 리포트 생성기"""
    
    HEADER_FILL = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    TOTAL_FILL = PatternFill(start_color="EBF3FF", end_color="EBF3FF", fill_type="solid")
    
    THIN_BORDER = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD')
    )
    
    HEADER_FONT = Font(name="맑은 고딕", size=10, bold=True, color="101828")
    BODY_FONT = Font(name="맑은 고딕", size=10, color="1D2939")
    FORMULA_FONT = Font(name="맑은 고딕", size=10, bold=True, color="0040C1")
    
    CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
    LEFT_ALIGN = Alignment(horizontal='left', vertical='center')
    RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')
    FEEDBACK_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    CORRECT_FILL = PatternFill(start_color="ECFDF3", end_color="ECFDF3", fill_type="solid")
    PARTIAL_FILL = PatternFill(start_color="FFFAEB", end_color="FFFAEB", fill_type="solid")
    WRONG_FILL = PatternFill(start_color="FEF3F2", end_color="FEF3F2", fill_type="solid")
    CONFIRMED_FILL = PatternFill(start_color="EFF8FF", end_color="EFF8FF", fill_type="solid")

    @classmethod
    def clean_sheet_name(cls, name: str) -> str:
        """엑셀 시트명에 사용할 수 없는 특수문자 제거 및 길이 제한(31자) 적용"""
        cleaned = re.sub(r'[\\/*?:\[\]]', '_', name.strip())
        return cleaned[:31] if cleaned else "기본학과"

    @classmethod
    def generate_excel_report(cls, criteria_data: Dict[str, Any], graded_students: List[Dict[str, Any]], include_detail_sheet: bool = True) -> bytes:
        """
        채점 결과 데이터를 구글 스프레드시트 규격에 맞춰 학과별 시트가 자동 분할된 엑셀 바이너리로 변환합니다.
        """
        wb = openpyxl.Workbook()
        default_sheet = wb.active  # 기본 생성 시트 보존
        
        # 1. 학생들을 학과별로 그룹화
        dept_groups: Dict[str, List[Dict[str, Any]]] = {}
        for student in graded_students:
            dept = student.get("department", "일반").strip() or "일반"
            if dept not in dept_groups:
                dept_groups[dept] = []
            dept_groups[dept].append(student)
            
        items = criteria_data.get("items", [])
        item_keys = [item["item_no"] for item in items]
        correct_answers = {item["item_no"]: item.get("correct_answer", "") for item in items}
        total_possible_score = sum([float(item.get("score", 0)) for item in items])
        
        # 학생이 0명인 경우에도 빈 양식 시트 1개 보장
        if not dept_groups:
            dept_groups["성적표_양식"] = []
            
        created_sheets = []
        # 2. 각 학과별로 워크시트 생성
        for dept_name, students in dept_groups.items():
            sheet_title = cls.clean_sheet_name(dept_name)
            ws = wb.create_sheet(title=sheet_title)
            created_sheets.append(ws)
            
            # --- 구글 스프레드시트 3행 헤더 구성 ---
            # 1행: A1="", B1="문항", C1~=문항번호들, 끝-1="득점", 끝="순위"
            row1 = ["", "문항"] + item_keys + ["득점", "순위"]
            ws.append(row1)
            
            # 2행: A2="", B2="정답", C2~=정답목록, 끝-1=총배점, 끝=""
            row2 = ["", "정답"] + [correct_answers.get(k, "") for k in item_keys] + [total_possible_score, ""]
            ws.append(row2)
            
            # 3행: A3="학번", B3="성명", C3~=""
            row3 = ["학번", "성명"] + ["" for _ in item_keys] + ["", ""]
            ws.append(row3)
            
            # 열 위치 계산
            start_item_col = 3  # C열
            end_item_col = 2 + len(item_keys)  # 마지막 문항 열
            score_col = end_item_col + 1  # 득점 열
            rank_col = score_col + 1      # 순위 열
            
            score_col_letter = get_column_letter(score_col)
            rank_col_letter = get_column_letter(rank_col)
            start_item_col_letter = get_column_letter(start_item_col)
            end_item_col_letter = get_column_letter(end_item_col)
            
            # 4행부터 학생별 점수 및 엑셀 수식 작성
            start_student_row = 4
            end_student_row = start_student_row + len(students) - 1
            
            for idx, student in enumerate(students):
                current_row = start_student_row + idx
                student_id = student.get("student_id", "")
                name = student.get("name", "")
                results = student.get("results", {})
                
                # 문항별 취득 점수
                item_scores = [results.get(k, {}).get("score", 0.0) for k in item_keys]
                
                # 득점 수식 (=SUM(C4:AG4))
                sum_formula = f"=SUM({start_item_col_letter}{current_row}:{end_item_col_letter}{current_row})"
                # 순위 수식 (=RANK(AH4, $AH$4:$AH$30))
                rank_formula = f"=RANK({score_col_letter}{current_row}, ${score_col_letter}${start_student_row}:${score_col_letter}${end_student_row})"
                
                row_data = [student_id, name] + item_scores + [sum_formula, rank_formula]
                ws.append(row_data)

            # 스타일링 적용 (폰트, 테두리, 배경색, 정렬)
            max_col = rank_col
            max_row = max(end_student_row, 3)
            
            # 1행~3행 (헤더 영역)
            for r in range(1, 4):
                for c in range(1, max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.border = cls.THIN_BORDER
                    cell.alignment = cls.CENTER_ALIGN
                    cell.font = cls.HEADER_FONT
                    if r == 1:
                        cell.fill = cls.HEADER_FILL
                    elif r == 2:
                        cell.fill = cls.SUBHEADER_FILL
                    elif r == 3:
                        cell.fill = cls.HEADER_FILL
            
            # 득점/순위 헤더 강조
            ws.cell(row=1, column=score_col).fill = cls.TOTAL_FILL
            ws.cell(row=1, column=rank_col).fill = cls.TOTAL_FILL
            
            # 학생 데이터 행 스타일
            for r in range(start_student_row, max_row + 1):
                for c in range(1, max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.border = cls.THIN_BORDER
                    cell.font = cls.BODY_FONT
                    
                    if c in [1, 2]:  # 학번, 성명
                        cell.alignment = cls.CENTER_ALIGN
                    elif c == score_col:  # 득점 (수식)
                        cell.alignment = cls.RIGHT_ALIGN
                        cell.font = cls.FORMULA_FONT
                        cell.fill = cls.TOTAL_FILL
                    elif c == rank_col:   # 순위 (수식)
                        cell.alignment = cls.CENTER_ALIGN
                        cell.font = cls.FORMULA_FONT
                    else:  # 문항별 점수
                        cell.alignment = cls.RIGHT_ALIGN
                        
            # 열 너비 자동 조정
            for col in range(1, max_col + 1):
                col_letter = get_column_letter(col)
                if col == 1:
                    ws.column_dimensions[col_letter].width = 14  # 학번
                elif col == 2:
                    ws.column_dimensions[col_letter].width = 12  # 성명
                elif col in [score_col, rank_col]:
                    ws.column_dimensions[col_letter].width = 10  # 득점, 순위
                else:
                    ws.column_dimensions[col_letter].width = 8   # 문항들
                    
            # 옵션: 상세 답안/정오표 시트도 함께 추가
            if include_detail_sheet:
                detail_title = cls.clean_sheet_name(f"{dept_name}_상세정오표")
                ws_detail = wb.create_sheet(title=detail_title)
                
                # 상세 시트 헤더: 문항별 [학생답안, 정오, 득점]
                detail_row1 = ["학번", "성명"]
                detail_row2 = ["", ""]
                for k in item_keys:
                    detail_row1.extend([k, "", ""])
                    detail_row2.extend(["답안", "정오", "득점"])
                detail_row1.extend(["총점"])
                detail_row2.extend(["합계"])
                
                ws_detail.append(detail_row1)
                ws_detail.append(detail_row2)
                
                for student in students:
                    student_id = student.get("student_id", "")
                    name = student.get("name", "")
                    results = student.get("results", {})
                    
                    row = [student_id, name]
                    tot = 0.0
                    for k in item_keys:
                        res = results.get(k, {})
                        row.extend([res.get("answer", ""), res.get("status", "X"), res.get("score", 0.0)])
                        tot += float(res.get("score", 0.0))
                    row.append(tot)
                    ws_detail.append(row)
                    
                # 상세 시트 서식 적용
                for r in range(1, ws_detail.max_row + 1):
                    for c in range(1, ws_detail.max_column + 1):
                        cell = ws_detail.cell(row=r, column=c)
                        cell.border = cls.THIN_BORDER
                        cell.alignment = cls.CENTER_ALIGN
                        if r in [1, 2]:
                            cell.fill = cls.HEADER_FILL
                            cell.font = cls.HEADER_FONT
                        else:
                            cell.font = cls.BODY_FONT

        # 불필요한 초기 빈 시트 제거 및 활성 시트 보장
        if created_sheets and default_sheet in wb.worksheets:
            wb.remove(default_sheet)
        if len(wb.sheetnames) > 0:
            wb.active = 0

        # 엑셀 바이너리 스트림 반환
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @classmethod
    def generate_student_feedback_report(cls, student: Dict[str, Any], criteria_data: Dict[str, Any]) -> bytes:
        """
        특정 학생 1명의 문항별 답안, 기준 정답, 정오, 득점, 상세 채점 피드백 및 수기 확정 내역을 담은 엑셀 리포트 생성
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        s_name = student.get("name", "학생")
        s_id = student.get("student_id", "")
        s_dept = student.get("department", "일반")
        ws.title = cls.clean_sheet_name(f"{s_name}_피드백")

        items = criteria_data.get("items", [])
        crit_map = {item["item_no"]: item for item in items}
        results = student.get("results", {})
        total_possible = sum([float(it.get("score", 0)) for it in items])
        student_total = float(student.get("total_score", 0.0))

        # 1. 상단 타이틀
        title_font = Font(name="맑은 고딕", size=14, bold=True, color="101828")
        ws.merge_cells("A1:I1")
        title_cell = ws.cell(row=1, column=1, value=f"📋 [{s_name}] 학생 문항별 채점 피드백 & 정오표")
        title_cell.font = title_font
        title_cell.alignment = cls.CENTER_ALIGN
        ws.row_dimensions[1].height = 30

        # 2. 학생 기본 정보 카드 (2행~3행)
        meta_headers = ["학과", "학번", "성명", "취득 총점", "총 배점", "성취율(%)", "수기 확정 문항"]
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 22

        confirmed_count = sum(1 for r in results.values() if r.get("confirmed_by_human"))
        rate = (student_total / total_possible * 100) if total_possible > 0 else 0.0

        meta_values = [
            s_dept,
            s_id,
            s_name,
            f"{student_total:.1f}점",
            f"{total_possible:.1f}점",
            f"{rate:.1f}%",
            f"{confirmed_count}건"
        ]

        for c_idx, (h, v) in enumerate(zip(meta_headers, meta_values), start=1):
            h_cell = ws.cell(row=2, column=c_idx, value=h)
            h_cell.font = cls.HEADER_FONT
            h_cell.fill = cls.SUBHEADER_FILL
            h_cell.alignment = cls.CENTER_ALIGN
            h_cell.border = cls.THIN_BORDER

            v_cell = ws.cell(row=3, column=c_idx, value=v)
            v_cell.font = Font(name="맑은 고딕", size=10, bold=(c_idx in [4, 6]))
            v_cell.alignment = cls.CENTER_ALIGN
            v_cell.border = cls.THIN_BORDER
            if c_idx == 4:
                v_cell.fill = cls.TOTAL_FILL

        # 3. 문항별 피드백 테이블 헤더 (5행)
        start_table_row = 5
        ws.row_dimensions[start_table_row].height = 24
        headers = ["문항", "유형", "배점", "기준 정답", "학생 답안", "정오", "득점", "채점 피드백", "판정 구분"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=start_table_row, column=c_idx, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.CENTER_ALIGN
            cell.border = cls.THIN_BORDER

        # 4. 문항 데이터 작성 (6행부터)
        cur_row = start_table_row + 1
        for item in items:
            k = item["item_no"]
            crit = crit_map.get(k, {})
            res = results.get(k, {})
            
            q_type = crit.get("type", "단답형")
            max_sc = float(crit.get("score", 0.0))
            correct_ans = str(crit.get("correct_answer", ""))
            student_ans = str(res.get("answer", ""))
            status = str(res.get("status", "X")).upper()
            score = float(res.get("score", 0.0))
            feedback = str(res.get("feedback", ""))
            is_confirmed = res.get("confirmed_by_human", False)
            judge_type = "수기 확정 ✏️" if is_confirmed else "AI 자동채점"

            row_vals = [k, q_type, max_sc, correct_ans, student_ans, status, score, feedback, judge_type]
            ws.row_dimensions[cur_row].height = 22

            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=cur_row, column=c_idx, value=val)
                cell.font = cls.BODY_FONT
                cell.border = cls.THIN_BORDER

                # 열별 정렬 및 스타일
                if c_idx in [1, 2, 6, 9]:
                    cell.alignment = cls.CENTER_ALIGN
                elif c_idx in [3, 7]:
                    cell.alignment = cls.RIGHT_ALIGN
                elif c_idx == 8:
                    cell.alignment = cls.FEEDBACK_ALIGN
                else:
                    cell.alignment = cls.LEFT_ALIGN

                # 정오 상태별 배경색 강조
                if c_idx == 6:
                    if status == "O":
                        cell.fill = cls.CORRECT_FILL
                        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="027A48")
                    elif status == "△":
                        cell.fill = cls.PARTIAL_FILL
                        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="B54708")
                    else:
                        cell.fill = cls.WRONG_FILL
                        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="B42318")

                # 수기 확정 구분 강조
                if c_idx == 9 and is_confirmed:
                    cell.fill = cls.CONFIRMED_FILL
                    cell.font = Font(name="맑은 고딕", size=10, bold=True, color="175CD3")

            cur_row += 1

        # 5. 합계 행
        ws.row_dimensions[cur_row].height = 24
        total_label_cell = ws.cell(row=cur_row, column=1, value="합계")
        total_label_cell.font = cls.HEADER_FONT
        total_label_cell.alignment = cls.CENTER_ALIGN
        total_label_cell.fill = cls.TOTAL_FILL
        total_label_cell.border = cls.THIN_BORDER
        ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=2)

        # 배점 합계
        tot_max_cell = ws.cell(row=cur_row, column=3, value=f"=SUM(C{start_table_row+1}:C{cur_row-1})")
        tot_max_cell.font = cls.FORMULA_FONT
        tot_max_cell.alignment = cls.RIGHT_ALIGN
        tot_max_cell.fill = cls.TOTAL_FILL
        tot_max_cell.border = cls.THIN_BORDER

        # 빈 셀 채우기
        for c in range(4, 7):
            empty_cell = ws.cell(row=cur_row, column=c, value="")
            empty_cell.fill = cls.TOTAL_FILL
            empty_cell.border = cls.THIN_BORDER

        # 득점 합계
        tot_sc_cell = ws.cell(row=cur_row, column=7, value=f"=SUM(G{start_table_row+1}:G{cur_row-1})")
        tot_sc_cell.font = cls.FORMULA_FONT
        tot_sc_cell.alignment = cls.RIGHT_ALIGN
        tot_sc_cell.fill = cls.TOTAL_FILL
        tot_sc_cell.border = cls.THIN_BORDER

        for c in range(8, 10):
            empty_cell = ws.cell(row=cur_row, column=c, value="")
            empty_cell.fill = cls.TOTAL_FILL
            empty_cell.border = cls.THIN_BORDER

        # 열 너비 설정
        col_widths = {
            1: 10,  # 문항
            2: 12,  # 유형
            3: 8,   # 배점
            4: 25,  # 기준 정답
            5: 25,  # 학생 답안
            6: 8,   # 정오
            7: 8,   # 득점
            8: 45,  # 채점 피드백
            9: 14   # 판정 구분
        }
        for c_idx, w in col_widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = w

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @classmethod
    def generate_all_feedback_report(cls, criteria_data: Dict[str, Any], graded_students: List[Dict[str, Any]]) -> bytes:
        """
        전체 학생의 문항별 채점 피드백 및 정오표 종합 리포트 생성 (엑셀)
        - 시트 1: 전체 피드백 목록 (필터링 및 분석용 전체 문항 레코드)
        - 시트 2: 수기 확정 문항 목록 (채점자가 수기 조정한 문항만 별도 집계)
        """
        wb = openpyxl.Workbook()
        ws_all = wb.active
        ws_all.title = "전체_피드백_종합"

        items = criteria_data.get("items", [])
        crit_map = {item["item_no"]: item for item in items}

        headers = ["학과", "학번", "성명", "문항", "유형", "배점", "기준 정답", "학생 답안", "정오", "득점", "채점 피드백", "판정 구분"]
        ws_all.row_dimensions[1].height = 24
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_all.cell(row=1, column=c_idx, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.CENTER_ALIGN
            cell.border = cls.THIN_BORDER

        cur_row = 2
        confirmed_records = []

        for student in graded_students:
            s_dept = student.get("department", "일반")
            s_id = student.get("student_id", "")
            s_name = student.get("name", "")
            results = student.get("results", {})

            for item in items:
                k = item["item_no"]
                crit = crit_map.get(k, {})
                res = results.get(k, {})

                q_type = crit.get("type", "단답형")
                max_sc = float(crit.get("score", 0.0))
                correct_ans = str(crit.get("correct_answer", ""))
                student_ans = str(res.get("answer", ""))
                status = str(res.get("status", "X")).upper()
                score = float(res.get("score", 0.0))
                feedback = str(res.get("feedback", ""))
                is_confirmed = res.get("confirmed_by_human", False)
                judge_type = "수기 확정 ✏️" if is_confirmed else "AI 자동채점"

                row_vals = [s_dept, s_id, s_name, k, q_type, max_sc, correct_ans, student_ans, status, score, feedback, judge_type]
                ws_all.row_dimensions[cur_row].height = 22

                for c_idx, val in enumerate(row_vals, start=1):
                    cell = ws_all.cell(row=cur_row, column=c_idx, value=val)
                    cell.font = cls.BODY_FONT
                    cell.border = cls.THIN_BORDER

                    if c_idx in [1, 2, 3, 4, 5, 9, 12]:
                        cell.alignment = cls.CENTER_ALIGN
                    elif c_idx in [6, 10]:
                        cell.alignment = cls.RIGHT_ALIGN
                    elif c_idx == 11:
                        cell.alignment = cls.FEEDBACK_ALIGN
                    else:
                        cell.alignment = cls.LEFT_ALIGN

                    if c_idx == 9:
                        if status == "O":
                            cell.fill = cls.CORRECT_FILL
                            cell.font = Font(name="맑은 고딕", size=10, bold=True, color="027A48")
                        elif status == "△":
                            cell.fill = cls.PARTIAL_FILL
                            cell.font = Font(name="맑은 고딕", size=10, bold=True, color="B54708")
                        else:
                            cell.fill = cls.WRONG_FILL
                            cell.font = Font(name="맑은 고딕", size=10, bold=True, color="B42318")

                    if c_idx == 12 and is_confirmed:
                        cell.fill = cls.CONFIRMED_FILL
                        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="175CD3")

                if is_confirmed:
                    confirmed_records.append(row_vals)

                cur_row += 1

        # 열 너비 설정
        all_col_widths = {
            1: 12,  # 학과
            2: 12,  # 학번
            3: 12,  # 성명
            4: 10,  # 문항
            5: 12,  # 유형
            6: 8,   # 배점
            7: 22,  # 기준 정답
            8: 22,  # 학생 답안
            9: 8,   # 정오
            10: 8,  # 득점
            11: 45, # 채점 피드백
            12: 14  # 판정 구분
        }
        for c_idx, w in all_col_widths.items():
            ws_all.column_dimensions[get_column_letter(c_idx)].width = w

        # 시트 2: 수기 확정 검토 목록 시트
        ws_conf = wb.create_sheet(title="수기_확정_검토내역")
        ws_conf.row_dimensions[1].height = 24
        for c_idx, h in enumerate(headers, start=1):
            cell = ws_conf.cell(row=1, column=c_idx, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.CONFIRMED_FILL
            cell.alignment = cls.CENTER_ALIGN
            cell.border = cls.THIN_BORDER

        for r_idx, r_vals in enumerate(confirmed_records, start=2):
            ws_conf.row_dimensions[r_idx].height = 22
            for c_idx, val in enumerate(r_vals, start=1):
                cell = ws_conf.cell(row=r_idx, column=c_idx, value=val)
                cell.font = cls.BODY_FONT
                cell.border = cls.THIN_BORDER
                if c_idx in [1, 2, 3, 4, 5, 9, 12]:
                    cell.alignment = cls.CENTER_ALIGN
                elif c_idx in [6, 10]:
                    cell.alignment = cls.RIGHT_ALIGN
                elif c_idx == 11:
                    cell.alignment = cls.FEEDBACK_ALIGN
                else:
                    cell.alignment = cls.LEFT_ALIGN

        for c_idx, w in all_col_widths.items():
            ws_conf.column_dimensions[get_column_letter(c_idx)].width = w

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

